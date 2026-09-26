"""Шорт-лист в Excel: «Кого» — отмеченные «Позвать» или все из «Подходят»; «Для кого» —
для себя (с контактами и внутренними пометками) или для клиента (без телефона, почты
и пометок). Одна строка — один кандидат, шапка закреплена, на колонках автофильтр."""

from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from sqlalchemy import select
from sqlalchemy.orm import Session

from app import evaluate as ev
from app.models import Candidate, Match, Vacancy

TRACE = "Оценки подготовил ИИ-агент, список проверил рекрутер"
HEAD = 5  # строка с названиями колонок
COLUMNS = [
    ("ФИО", 28),
    ("Категория", 18),
    ("Оценка", 8),
    ("Почему подходит", 60),
    ("Что настораживает", 50),
    ("Вопросы на созвон", 50),
    ("Город", 14),
    ("Стаж, лет", 9),
    ("Дата резюме", 12),
]
OWN = [("Телефон", 16), ("Почта", 26), ("Решение", 14), ("Пометки", 40)]
DECISION = {"invite": "Позвать", "maybe": "Под вопросом"}


def chosen(session: Session, v: Vacancy, who: str) -> list[tuple[Match, Candidate]]:
    rows = session.execute(
        select(Match, Candidate)
        .join(Candidate, Candidate.id == Match.candidate_id)
        .where(Match.vacancy_id == v.id, Match.decision.is_distinct_from("reject"))
    ).all()
    if who == "invite":
        rows = [(m, c) for m, c in rows if m.decision == "invite"]
    else:
        rows = [(m, c) for m, c in rows if ev.category(m) == ev.FIT]
    return sorted(rows, key=lambda r: (bool(r[1].stale), -(r[0].score or 0)))


def _reasons(m: Match) -> str:
    return "\n".join(
        f"{r['name']}: {r['quote']}" if r.get("quote") else r["name"] for r in m.reasons
    )


def _concerns(m: Match) -> str:
    return "\n".join(
        f"{x['text']} ({'в резюме: ' + x['quote'] if x.get('quote') else 'в резюме не сказано'})"
        for x in m.concerns
    )


def shortlist(session: Session, v: Vacancy, who: str, audience: str) -> bytes:
    own = audience == "self"
    columns = COLUMNS + (OWN if own else [])
    wb = Workbook()
    ws = wb.active
    ws.title = "Шорт-лист"
    ws["A1"] = v.title or "Вакансия"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Дата: {date.today():%d.%m.%Y}"
    ws["A3"] = TRACE
    ws["A3"].font = Font(italic=True)
    for col, (name, width) in enumerate(columns, start=1):
        cell = ws.cell(HEAD, col, name)
        cell.font = Font(bold=True)
        ws.column_dimensions[cell.column_letter].width = width
    row = HEAD
    for m, c in chosen(session, v, who):
        row += 1
        p = c.parsed or {}
        values = [
            c.full_name or "Без имени",
            ev.CATEGORIES[ev.category(m)],
            m.score,
            _reasons(m),
            _concerns(m),
            "\n".join(m.questions),
            c.city or p.get("city") or "",
            p.get("total_years"),
            c.resume_date,
        ]
        if own:
            notes = [f"не подходит: {m.decision_reason}"] if m.decision_reason else []
            if c.stale:
                notes.append("резюме старше полутора лет")
            values += [
                ", ".join(c.phones),
                ", ".join(c.emails),
                DECISION.get(m.decision or "", ""),
                "; ".join(notes),
            ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row, col, value)
            if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
                cell.data_type = "s"  # текст из резюме не становится формулой
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if isinstance(value, date):
                cell.number_format = "DD.MM.YYYY"
    ws.freeze_panes = ws.cell(HEAD + 1, 1)
    ws.auto_filter.ref = f"A{HEAD}:{ws.cell(HEAD, len(columns)).column_letter}{max(row, HEAD)}"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
