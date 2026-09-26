"""Чтение выгрузок (CSV, XLSX) и файлов резюме (PDF, DOCX, TXT, ZIP)."""

import csv
import io
import zipfile
from functools import lru_cache
from pathlib import Path

from openpyxl import load_workbook

DOC_SUFFIXES = {".pdf", ".docx", ".txt"}
TABLE_SUFFIXES = {".csv", ".xlsx"}


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _cell(value):
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    return value  # даты остаются датами, строки — строками


def read_table(path: Path) -> tuple[list[str], list[list]]:
    """Заголовки и строки таблицы; пустые строки отбрасываются."""
    if path.suffix.lower() == ".xlsx":
        wb = load_workbook(path, read_only=True, data_only=True)
        raw = [[_cell(v) for v in row] for row in wb.active.iter_rows(values_only=True)]
        wb.close()
    else:
        text = _decode(path.read_bytes())
        # Разделитель — по строке заголовков: Sniffer путается на многострочных ячейках.
        header = text.split("\n", 1)[0]
        delimiter = max(",;\t", key=header.count)
        raw = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    rows = [r for r in raw if any(str(v or "").strip() for v in r)]
    if not rows:
        return [], []
    headers = [str(h or "").strip() or f"Колонка {i + 1}" for i, h in enumerate(rows[0])]
    return headers, rows[1:]


def _zip_name(info: zipfile.ZipInfo) -> str:
    # Архивы из Windows хранят русские имена в cp866 без флага UTF-8.
    if info.flag_bits & 0x800:
        return info.filename
    try:
        return info.filename.encode("cp437").decode("cp866")
    except (UnicodeEncodeError, UnicodeDecodeError):
        return info.filename


def extract_zip(archive: Path, dest: Path) -> None:
    with zipfile.ZipFile(archive) as z:
        for info in z.infolist():
            if info.is_dir():
                continue
            name = Path(_zip_name(info)).name
            if Path(name).suffix.lower() in DOC_SUFFIXES and not name.startswith("."):
                target = dest / name
                n = 1
                while target.exists():
                    target = dest / f"{Path(name).stem}_{n}{Path(name).suffix}"
                    n += 1
                target.write_bytes(z.read(info))


def list_documents(folder: Path) -> list[Path]:
    if not folder.exists():
        return []
    files = [
        p
        for p in folder.rglob("*")
        if p.is_file() and p.suffix.lower() in DOC_SUFFIXES and not p.name.startswith(".")
    ]
    return sorted(files)


@lru_cache(maxsize=1)
def _converters() -> dict:
    # Конвертеры markitdown зовём напрямую по расширению: общий MarkItDown() на каждый
    # файл запускает нейросеть magika для угадывания типа (~0,5 с на файл).
    from markitdown.converters import DocxConverter, PdfConverter, PlainTextConverter

    return {".docx": DocxConverter(), ".pdf": PdfConverter(), ".txt": PlainTextConverter()}


def document_text(path: Path) -> str:
    from markitdown import StreamInfo

    suffix = path.suffix.lower()
    info = StreamInfo(extension=suffix, filename=path.name)
    if suffix == ".txt":
        info = StreamInfo(extension=suffix, filename=path.name, charset=_charset(path))
    with path.open("rb") as stream:
        return _converters()[suffix].convert(stream, info).markdown.strip()


def _charset(path: Path) -> str:
    try:
        path.read_bytes().decode("utf-8")
        return "utf-8-sig"
    except UnicodeDecodeError:
        return "cp1251"
