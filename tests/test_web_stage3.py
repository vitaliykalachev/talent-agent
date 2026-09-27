"""Экраны этапа 3 на моке: вакансии, портрет, результат, решения, «Неверно», экспорт."""

import re
from datetime import date
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select
from test_evaluate import answer, base, matches, mock, run, standard  # noqa: F401 — фикстуры

from app import evaluate as ev
from app.main import create_app
from app.models import Feedback, Job, Vacancy

BANNED = ("провайдер", "токен", "эмбеддинг", "парсинг", "cron", "prompt", "llm")


@pytest.fixture
def web(session):
    return TestClient(create_app())


def page(web, url, **kw) -> str:
    r = web.get(url, **kw)
    assert r.status_code == 200, url
    text = re.sub(r"<[^>]+>", " ", r.text).lower()  # слова на экране, без разметки
    for word in BANNED:
        assert word not in text, (url, word)
    return r.text


@pytest.mark.parametrize("url", ["/vacancies", "/vacancies/new"])
def test_screens_on_empty_base(client, url):
    r = client.get(url)
    assert r.status_code == 200
    assert "Новая вакансия" in r.text


def test_new_vacancy_parsed_into_portrait_with_offers(web, session, base, mock):  # noqa: F811
    v, p = base
    mock(
        [
            {
                "match": "Нужен начальник литейного цеха",
                "schema": "VacancyDraft",
                "response": {
                    "title": "Начальник литейного цеха",
                    "requirements": [
                        {"name": "Руководил цехом литья от 3 лет", "kind": "must"},
                        {"name": "Литьё под давлением", "kind": "must"},
                        {
                            "name": "Хорошие навыки управления",
                            "kind": "nice",
                            "flag": "vague",
                            "suggestion": "Руководил коллективом от 30 человек",
                        },
                        {"name": "Мужчина до 45 лет", "kind": "nice"},
                        {"name": "Только продажи", "kind": "avoid"},
                    ],
                },
            }
        ]
    )
    r = web.post(
        "/vacancies/new",
        data={
            "description": "Нужен начальник литейного цеха, литьё под давлением.",
            "city": "",
            "top_n": "40",
        },
        follow_redirects=False,
    )
    assert r.status_code == 303
    url = r.headers["location"]
    text = page(web, url)
    for words in (
        "Портрет кандидата",
        "Обязательно",
        "Желательно",
        "Чего точно не надо",
        "Руководил цехом литья от 3 лет",
        "Расплывчато.",
        "Заменить на «Руководил коллективом от 30 человек»",
        "Так писать нельзя.",
        "Оценить 5 для проверки",
        "Оценить 40",
        "₽",
        "займёт",
        "Агент запомнил",
    ):
        assert words in text, words
    vid = int(url.rsplit("/", 1)[1])
    fresh = session.get(Vacancy, vid)
    rid = next(r["id"] for r in fresh.requirements if r["flag"] == "vague")
    web.post(f"{url}/requirements", data={"action": "accept", "rid": rid})
    web.post(f"{url}/requirements", data={"action": "add", "name": "SAP PP", "kind": "nice"})
    session.expire_all()
    assert session.get(Vacancy, vid).requirements_version == 3
    for i in range(7):  # больше десяти требований — предупреждение
        web.post(
            f"{url}/requirements", data={"action": "add", "name": f"Навык {i}", "kind": "nice"}
        )
    assert "Требований больше 10" in page(web, url)


def test_short_description_is_asked_again(web, session):
    r = web.post("/vacancies/new", data={"description": "Нужен"})
    assert r.status_code == 400 and "хотя бы парой предложений" in r.text


def test_results_groups_card_and_resume_link(web, session, base, mock):  # noqa: F811
    v, p = base
    mock(standard(v, p))
    web.post(f"/vacancies/{v.id}/evaluate", data={"scope": "trial"})
    from app.jobs import run_pending

    run_pending()
    text = page(web, f"/vacancies/{v.id}/results")
    assert "Подходят" in text and "Скорее не подходят" in text
    assert '<details class="group unfit">' in text  # свёрнута, но в списке
    assert '<details class="group fit" open>' in text
    assert "обязательные 2 из 2 · желательные 2 из 2" in text
    assert "Почему подходит" in text and "Что настораживает" not in text.split("Орлова")[0]
    # аудит №13 и №14: после первой оценки «Новый» ни у кого; оговорка — строка без кнопки
    assert "Новый" not in text and "Понятно" not in text and "может ошибиться" in text
    assert "Вопросы на первый созвон" in text and "Какой был объём выпуска?" in text
    assert "в резюме не сказано" in text  # у Орловой по трём требованиям нет данных
    link = f"/candidates/{p['Громов'].id}?v={v.id}#e-r1"
    assert link in text
    card = page(web, link.split("#")[0])
    assert '<mark id="e-r1">' in card and "Подсвечены строки" in card
    assert "Оценки по вакансиям" in card and "Начальник цеха литья" in card


def test_decisions_hide_rejected_and_survive_rerun(web, session, base, mock):  # noqa: F811
    v, p = base
    mock(standard(v, p))
    run(session, v)
    sid = p["Сидоров"].id
    web.post(
        f"/vacancies/{v.id}/decision",
        data={"candidate_id": sid, "decision": "reject", "reason": "Нет нужного опыта"},
    )
    web.post(
        f"/vacancies/{v.id}/decision", data={"candidate_id": p["Громов"].id, "decision": "invite"}
    )
    text = page(web, f"/vacancies/{v.id}/results")
    assert "Показать скрытых (1)" in text and f'id="c{sid}"' not in text
    hidden = page(web, f"/vacancies/{v.id}/results", params={"hidden": 1})
    assert f'id="c{sid}"' in hidden and "Не подходит: нет нужного опыта" in hidden
    ev.add_requirement(session, v, "SAP PP", "nice")
    run(session, v)
    got = matches(session, v)
    assert got["Сидоров"].decision == "reject" and got["Громов"].decision == "invite"
    text = page(web, f"/vacancies/{v.id}/results")
    assert "Оценка изменилась: 100 →" in text and "вы изменили требования вакансии" in text


def test_wrong_reason_writes_feedback_with_undo_and_memory(web, session, base, mock):  # noqa: F811
    v, p = base
    mock(standard(v, p))
    run(session, v)
    r = web.post(
        f"/vacancies/{v.id}/feedback",
        data={
            "candidate_id": p["Громов"].id,
            "target": "reason",
            "text": "Бережливое производство",
            "kind": "not_in_resume",
            "scope": "vacancy",
        },
        follow_redirects=False,
    )
    f = session.scalar(select(Feedback))
    assert (f.kind, f.scope, f.vacancy_id, f.target) == ("not_in_resume", "vacancy", v.id, "reason")
    text = page(web, r.headers["location"])
    assert (
        "Агент учтёт это в вакансии «Начальник цеха литья»: при следующей оценке пересчитает 1"
        in text
    )
    assert "Отменить" in text
    memory = page(web, f"/vacancies/{v.id}")
    assert "«Бережливое производство» — этого нет в резюме" in memory
    web.post(f"/feedback/{f.id}/delete", data={"back": f"/vacancies/{v.id}/results"})
    assert session.scalar(select(Feedback)) is None

    web.post(
        f"/vacancies/{v.id}/feedback",
        data={"text": "ISO 9001", "kind": "irrelevant", "scope": "all", "target": "reason"},
    )
    session.expire_all()
    f = session.scalar(select(Feedback))
    assert f.scope == "all" and f.vacancy_id is None


def test_stale_resume_lower_in_group_with_note(web, session, base, mock):  # noqa: F811
    v, p = base
    items = standard(v, p)
    items[1] = answer(v, p["Орлова"], {rid: ch for rid, ch in _all_met().items()})
    mock(items)
    p["Громов"].stale, p["Громов"].resume_date = True, date(2024, 3, 12)
    p["Орлова"].raw_text += "\nруководила цехом литья, литьё под давлением, ISO 9001"
    session.commit()
    run(session, v)
    text = page(web, f"/vacancies/{v.id}/results")
    fit = text.split("<h2>Подходят")[1].split("<h2>")[0]
    assert fit.index(f'id="c{p["Орлова"].id}"') < fit.index(f'id="c{p["Громов"].id}"')
    assert "Резюме от 12\u00a0мар 2024 — старше полутора лет, уточните на созвоне" in text


def _all_met():
    return {
        "r1": ("met", {"__lines__": "Мастер участка литья"}),
        "r2": ("met", {"__lines__": "литьё алюминия"}),
        "r3": ("met", {"__lines__": "5S"}),
        "r4": ("met", {"__lines__": "ISO 9001"}),
        "r5": ("not_met", None),
    }


def test_double_click_evaluate_starts_one_job(web, session, base, mock):  # noqa: F811
    v, p = base
    mock(standard(v, p))
    for _ in range(2):
        web.post(f"/vacancies/{v.id}/evaluate", data={"scope": "all"})
    assert len(list(session.scalars(select(Job).where(Job.kind == "evaluate")))) == 1


def _cells(data: bytes) -> tuple:
    ws = load_workbook(BytesIO(data)).active
    rows = [[c.value for c in row] for row in ws.iter_rows()]
    return ws, rows


def test_export_for_client_has_no_contacts_and_for_self_has(web, session, base, mock):  # noqa: F811
    v, p = base
    mock(standard(v, p))
    run(session, v)
    g = p["Громов"]
    g.phones, g.emails = ["+79175551234"], ["gromov.ilya@mail.ru"]
    session.commit()
    web.post(f"/vacancies/{v.id}/decision", data={"candidate_id": g.id, "decision": "invite"})

    r = web.get(f"/vacancies/{v.id}/export", params={"who": "invite", "audience": "client"})
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    ws, rows = _cells(r.content)
    assert rows[0][0] == "Начальник цеха литья" and rows[1][0].startswith("Дата: ")
    assert rows[2][0] == "Оценки подготовил ИИ-агент, список проверил рекрутер"
    head = rows[4]
    assert head[:9] == [
        "ФИО",
        "Категория",
        "Оценка",
        "Почему подходит",
        "Что настораживает",
        "Вопросы на созвон",
        "Город",
        "Стаж, лет",
        "Дата резюме",
    ]
    assert "Телефон" not in head and "Решение" not in head and "Пометки" not in head
    assert len(rows) == 6 and rows[5][0] == "Громов Илья Сергеевич" and rows[5][1] == "Подходят"
    flat = " ".join(str(x) for row in rows for x in row if x is not None)
    assert "555" not in re.sub(r"\D", "", flat)[-40:] and "@" not in flat
    assert "917" not in flat and "gromov" not in flat
    assert ws.freeze_panes == "A6" and ws.auto_filter.ref.startswith("A5:")

    r = web.get(f"/vacancies/{v.id}/export", params={"who": "fit", "audience": "self"})
    ws, rows = _cells(r.content)
    assert rows[4][-4:] == ["Телефон", "Почта", "Решение", "Пометки"]
    assert rows[5][-4:-1] == ["+79175551234", "gromov.ilya@mail.ru", "Позвать"]
